#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sdmx_file_to_excel.py - Lettura ROBUSTA di dati SDMX (ISTAT) -> Excel.
=====================================================================
Riconosce DA SOLO il formato del contenuto e lo legge:
  - XML structure-specific  (Series con attributi)
  - XML generic             (SeriesKey/Value + Obs)
  - SDMX-JSON               (structure.dimensions + dataSets)  [etichette incluse]
  - SDMX-CSV / CSV ISTAT    (sia "etichettato" sia standard a codici)
Usabile sia su un FILE scaricato, sia sui BYTE di una risposta web.
Le etichette in italiano: arrivano gia' nei formati JSON e CSV-etichettato;
per l'XML a codici si possono decodificare passando il file "SDMX struttura".

USO (file):
  python sdmx_file_to_excel.py  DATI.(xml|json|csv)  [--struttura STRUTTURA.xml]
        [--layout long|istat] [--righe REF_AREA] [--outdir CARTELLA] [--lang it|en]

Dipendenze: pandas, openpyxl, lxml
"""
from __future__ import annotations
import argparse, datetime as _dt, io, json, os, re, sys
import pandas as pd
from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

VALUE_NAMES = ("OBS_VALUE", "Osservazione", "obsValue", "Value", "VALUE")
TIME_NAMES = ("TIME_PERIOD", "Periodo", "Tempo", "TIME")
META_CSV_COLS = {"DATAFLOW", "STRUCTURE", "STRUCTURE_ID", "STRUCTURE_NAME", "ACTION"}


class SdmxError(Exception):
    pass


# ===========================================================================
#  DECODIFICA ETICHETTE dal file STRUTTURA (opzionale, per XML a codici)
# ===========================================================================
def _name_it(elem, lang):
    by = {}
    for n in elem.xpath("./*[local-name()='Name']"):
        l = n.get("{http://www.w3.org/XML/1998/namespace}lang", "")
        if n.text:
            by[l] = n.text.strip()
    return by.get(lang) or by.get("en") or next(iter(by.values()), None)


def carica_struttura(path, lang):
    if not path:
        return {}, {}, {}
    try:
        root = etree.parse(path).getroot()
    except (OSError, etree.XMLSyntaxError) as exc:
        raise SdmxError(f"Impossibile leggere il file struttura: {exc}")
    codelists = {}
    for cl in root.xpath("//*[local-name()='Codelist']"):
        cid = cl.get("id")
        if not cid:
            continue
        codes = {}
        for c in cl.xpath("./*[local-name()='Code']"):
            k = c.get("id")
            if k is not None:
                codes[k] = _name_it(c, lang) or k
        codelists[cid] = codes
    concept_names = {}
    for cs in root.xpath("//*[local-name()='ConceptScheme']"):
        for con in cs.xpath("./*[local-name()='Concept']"):
            i = con.get("id")
            if i:
                concept_names[i] = _name_it(con, lang) or i
    dim2cl, dim_names = {}, {}
    for dim in root.xpath("//*[local-name()='DimensionList']/*[local-name()='Dimension'] | "
                          "//*[local-name()='DimensionList']/*[local-name()='TimeDimension']"):
        did = dim.get("id")
        if not did:
            continue
        ref = dim.xpath(".//*[local-name()='Enumeration']/*[local-name()='Ref']")
        if ref:
            dim2cl[did] = ref[0].get("id")
        cref = dim.xpath(".//*[local-name()='ConceptIdentity']/*[local-name()='Ref']")
        dim_names[did] = (concept_names.get(cref[0].get("id")) if cref else None) or did
    return dim2cl, codelists, dim_names


# ===========================================================================
#  PARSER XML (structure-specific + generic)
# ===========================================================================
def _iter_series(root):
    rows, dim_ids, seen = [], [], set()
    for el in root.iter():
        if etree.QName(el).localname != "Series":
            continue
        attribs = dict(el.attrib)
        if attribs:  # structure-specific
            for k in attribs:
                if k not in seen:
                    seen.add(k); dim_ids.append(k)
            for obs in el.xpath("./*[local-name()='Obs']"):
                t = next((obs.get(a) for a in ("TIME_PERIOD", "TIME") if obs.get(a) is not None), None)
                v = next((obs.get(a) for a in ("OBS_VALUE", "value") if obs.get(a) is not None), None)
                r = dict(attribs); r["TIME_PERIOD"] = t; r["__VAL"] = v
                rows.append(r)
        else:  # generic
            base = {}
            for kv in el.xpath("./*[local-name()='SeriesKey']/*[local-name()='Value']"):
                cid = kv.get("id") or kv.get("concept")
                if cid:
                    if cid not in seen:
                        seen.add(cid); dim_ids.append(cid)
                    base[cid] = kv.get("value")
            for obs in el.xpath("./*[local-name()='Obs']"):
                od = obs.xpath("./*[local-name()='ObsDimension']")
                ov = obs.xpath("./*[local-name()='ObsValue']")
                r = dict(base)
                r["TIME_PERIOD"] = od[0].get("value") if od else None
                r["__VAL"] = ov[0].get("value") if ov else None
                rows.append(r)
    return rows, dim_ids


def parse_xml(content, struttura, lang):
    try:
        root = etree.fromstring(content if isinstance(content, bytes) else content.encode())
    except etree.XMLSyntaxError as exc:
        raise SdmxError(f"XML non valido: {exc}")
    rows, dim_ids = _iter_series(root)
    if not rows:
        raise SdmxError("XML letto ma senza osservazioni (Series/Obs non trovate).")
    dim2cl, codelists, dim_names = carica_struttura(struttura, lang) if struttura else ({}, {}, {})
    df = pd.DataFrame(rows)
    df["VALORE"] = pd.to_numeric(df["__VAL"], errors="coerce")
    df = df.drop(columns=["__VAL"])
    rename = {}
    for did in dim_ids:
        if did not in df.columns:
            continue
        cl = codelists.get(dim2cl.get(did, ""))
        if cl:
            df[did] = df[did].map(lambda v: cl.get(v, v))
        rename[did] = dim_names.get(did, did)
    cols = [c for c in dim_ids if c in df.columns] + ["TIME_PERIOD", "VALORE"]
    df = df[cols].rename(columns=rename)
    tl = dim_names.get("TIME_PERIOD", "TIME_PERIOD")
    df = df.rename(columns={"TIME_PERIOD": tl})
    df = df.sort_values([c for c in df.columns if c != "VALORE"], kind="stable").reset_index(drop=True)
    return df, tl


# ===========================================================================
#  PARSER SDMX-JSON  (etichette gia' incluse nei 'name')
# ===========================================================================
def parse_jsonstat(js, lang, codici=False):
    """JSON-stat 2.0 (class='dataset'): id/size/dimension/value -> df lungo.
    E' il formato del databrowserhub ISTAT (coeweb).
    codici=True aggiunge, per ogni dimensione, la colonna '<ID>_CODE' col codice."""
    ids = js.get("id") or list(js.get("dimension", {}).keys())
    size = js.get("size") or [len(js["dimension"][d]["category"]["label"]) for d in ids]
    dims = js["dimension"]
    pos2code, labels, names = [], [], []
    for did in ids:
        cat = dims[did]["category"]
        idx = cat.get("index")
        if isinstance(idx, dict):
            arr = [None] * len(idx)
            for code, pos in idx.items():
                arr[pos] = code
        elif isinstance(idx, list):
            arr = list(idx)
        else:
            arr = list(cat.get("label", {}).keys())
        pos2code.append(arr)
        labels.append(cat.get("label", {}))
        names.append(dims[did].get("label") or did)
    n = len(ids)
    strides = [1] * n
    for k in range(n - 2, -1, -1):
        strides[k] = strides[k + 1] * size[k + 1]
    value = js.get("value", {})
    items = value.items() if isinstance(value, dict) else enumerate(value)
    tpos = ids.index("TIME_PERIOD") if "TIME_PERIOD" in ids else (n - 1)
    rows = []
    for key, val in items:
        if val is None or val == "":
            continue
        i = int(key)
        row = {}
        for k, did in enumerate(ids):
            pos = (i // strides[k]) % size[k]
            code = pos2code[k][pos] if pos < len(pos2code[k]) else str(pos)
            row[names[k]] = labels[k].get(code, code)
            if codici:
                row[did + "_CODE"] = code
        try:
            row["VALORE"] = float(val)
        except (TypeError, ValueError):
            row["VALORE"] = val
        rows.append(row)
    if not rows:
        raise SdmxError("JSON-stat senza osservazioni.")
    df = pd.DataFrame(rows)
    time_label = names[tpos]
    df = df.sort_values([c for c in df.columns if c != "VALORE"], kind="stable").reset_index(drop=True)
    return df, time_label


def parse_json(content, lang):
    try:
        if isinstance(content, bytes):
            content = content.decode("utf-8-sig")
        js = json.loads(content)
    except ValueError as exc:
        raise SdmxError(f"JSON non valido: {exc}")
    # JSON-stat 2.0 (databrowserhub / coeweb)
    if js.get("class") == "dataset" or ("value" in js and "dimension" in js):
        return parse_jsonstat(js, lang)
    root = js.get("data", js)
    if "structure" not in root or "dataSets" not in root:
        raise SdmxError("JSON non in formato SDMX (mancano 'structure'/'dataSets').")
    dims = root["structure"].get("dimensions", {})
    sdims = dims.get("series", []); odims = dims.get("observation", [])
    ds = root["dataSets"][0]
    series = ds.get("series", {})
    cols = [d.get("name") or d.get("id") for d in sdims]
    tl = (odims[0].get("name") or odims[0].get("id")) if odims else "TIME_PERIOD"
    rows = []
    for skey, sval in series.items():
        idx = [int(i) for i in skey.split(":")]
        base = {}
        for pos, d in enumerate(sdims):
            vals = d.get("values", [])
            base[cols[pos]] = (vals[idx[pos]].get("name") or vals[idx[pos]].get("id")) if (pos < len(idx) and idx[pos] < len(vals)) else None
        for okey, ov in sval.get("observations", {}).items():
            r = dict(base); t = int(okey.split(":")[0])
            r[tl] = (odims[0]["values"][t].get("name") or odims[0]["values"][t].get("id")) if (odims and t < len(odims[0]["values"])) else None
            r["VALORE"] = pd.to_numeric(pd.Series([ov[0] if ov else None]), errors="coerce")[0]
            rows.append(r)
    if not rows:
        raise SdmxError("JSON SDMX senza osservazioni.")
    df = pd.DataFrame(rows)
    df = df.sort_values([c for c in df.columns if c != "VALORE"], kind="stable").reset_index(drop=True)
    return df, tl


# ===========================================================================
#  PARSER CSV  (etichettato ISTAT oppure standard a codici)
# ===========================================================================
def parse_csv(content):
    import csv as _csv
    text = content.decode("utf-8-sig") if isinstance(content, bytes) else content
    first = text.splitlines()[0] if text else ""
    sep = ";" if first.count(";") > first.count(",") else ","
    raw = pd.read_csv(io.StringIO(text), sep=sep, dtype=str, keep_default_na=False, quoting=_csv.QUOTE_NONE)
    cols = list(raw.columns)
    value_col = next((c for c in cols if c in VALUE_NAMES or c.upper() == "OBS_VALUE"), None)
    if value_col is None:
        raise SdmxError("CSV: colonna valore non trovata. Colonne presenti: " + ", ".join(cols))
    time_col = next((c for c in cols if c in TIME_NAMES or c.upper() == "TIME_PERIOD"), None)
    if time_col is None:
        raise SdmxError("CSV: colonna TIME_PERIOD non trovata. Colonne: " + ", ".join(cols))
    is_code = lambda c: bool(re.fullmatch(r"[A-Z0-9_]+", c))
    before = [c for c in cols[:cols.index(time_col)] if c != value_col]
    labelled = len(before) >= 2 and not is_code(before[1])
    if labelled:
        dim_cols = before[1::2]                       # etichetta di ogni coppia (codice,label)
    else:
        dim_cols = [c for c in before if c.upper() not in META_CSV_COLS]
    def pulisci(v):
        if isinstance(v, str) and len(v) >= 2 and v.startswith("'") and v.endswith("'"):
            v = v[1:-1].replace('"\'', "'")
        return v
    out = raw[dim_cols + [time_col]].copy()
    for c in out.columns:
        out[c] = out[c].map(pulisci)
    out["VALORE"] = pd.to_numeric(raw[value_col], errors="coerce")
    out = out.rename(columns={time_col: "TIME_PERIOD"})
    out = out.sort_values([c for c in out.columns if c != "VALORE"], kind="stable").reset_index(drop=True)
    return out, "TIME_PERIOD"


# ===========================================================================
#  RICONOSCIMENTO AUTOMATICO DEL FORMATO
# ===========================================================================
def sniff_e_parse(content, struttura=None, lang="it"):
    """content: bytes o str. Riconosce XML/JSON/CSV e ritorna (df, time_label)."""
    if isinstance(content, str):
        content = content.encode("utf-8")
    c = content[3:] if content[:3] == b"\xef\xbb\xbf" else content   # togli BOM
    head = c.lstrip()[:1]
    if head == b"<":
        return parse_xml(content, struttura, lang)
    if head in (b"{", b"["):
        return parse_json(content, lang)
    return parse_csv(content)


# ===========================================================================
#  FILTRO PERIODO (anno/trimestre) sul DataFrame gia' letto
# ===========================================================================
_RE_ANNO = re.compile(r"(\d{4})")
_RE_TRIM = re.compile(r"[QqTt](?:rim(?:estre)?)?\s*[.\-]?\s*0?([1-4])")


def _periodo_range(s):
    """Interpreta un periodo ('2025', '2025-Q1', 'T1-2025', '2025-03', ...)
    come (anno, trim_min, trim_max). None se non riconoscibile."""
    s = str(s)
    ma = _RE_ANNO.search(s)
    if not ma:
        return None
    anno = int(ma.group(1))
    resto = s[:ma.start()] + " " + s[ma.end():]          # togli l'anno, cerca il resto
    mt = _RE_TRIM.search(resto)
    if mt:
        t = int(mt.group(1))
        return anno, t, t
    mm = re.fullmatch(r"\s*\d{4}[-/](0?[1-9]|1[0-2])\s*", s)   # mensile "2025-03"
    if mm:
        t = (int(mm.group(1)) - 1) // 3 + 1
        return anno, t, t
    mM = re.search(r"[Mm]\s*0?(1[0-2]|[1-9])\b", resto)        # mensile "M03"
    if mM:
        t = (int(mM.group(1)) - 1) // 3 + 1
        return anno, t, t
    return anno, 1, 4                                    # annuale: copre tutto l'anno


def filtra_periodo(df, time_label, da=None, a=None, auto=False):
    """Tiene solo le righe del periodo scelto. da/a = (anno, trimestre).
    auto=True: ultimo anno con dati + tutto l'anno precedente."""
    if time_label not in df.columns or df.empty:
        return df
    ranges = {v: _periodo_range(v) for v in df[time_label].dropna().unique()}
    validi = [r for r in ranges.values() if r]
    if not validi:
        return df           # formato periodo non riconosciuto: meglio non tagliare
    if auto:
        da, a = (max(r[0] for r in validi) - 1, 1), None

    def _keep(v):
        r = ranges.get(v)
        if r is None:
            return True
        anno, t1, t2 = r
        if da and (anno, t2) < tuple(da):
            return False
        if a and (anno, t1) > tuple(a):
            return False
        return True

    out = df[df[time_label].map(_keep)].reset_index(drop=True)
    if out.empty:
        raise SdmxError("Nessun dato nel periodo scelto (controlla anni/trimestri).")
    return out


# ===========================================================================
#  EXCEL
# ===========================================================================
def _st():
    return {"hf": Font(size=11, bold=True, color="FFFFFF"), "hfill": PatternFill("solid", fgColor="2E75B6"),
            "alt": PatternFill("solid", fgColor="F2F6FC"), "bd": Border(*[Side(style="thin", color="BFBFBF")] * 4),
            "ctr": Alignment(horizontal="center", vertical="center", wrap_text=True),
            "rgt": Alignment(horizontal="right", vertical="center"), "lft": Alignment(horizontal="left", vertical="center")}


def _safe(t):
    return (re.sub(r"[^\w\-]+", "_", t, flags=re.UNICODE).strip("_")[:80]) or "dati"


def _meta(wb, nome, df):
    wm = wb.create_sheet("Metadati")
    wm.cell(row=1, column=1, value="METADATI").font = Font(bold=True, size=13, color="1F4E78")
    for i, (k, v) in enumerate([("Origine", nome), ("Data conversione", _dt.datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
                                ("Numero osservazioni", len(df)), ("Colonne", ", ".join(map(str, df.columns)))], start=3):
        wm.cell(row=i, column=1, value=k).font = Font(bold=True)
        wm.cell(row=i, column=2, value=str(v)).alignment = Alignment(wrap_text=True, vertical="top")
    wm.column_dimensions["A"].width = 22; wm.column_dimensions["B"].width = 100


def export_long(df, nome, outdir, time_label):
    st = _st(); os.makedirs(outdir, exist_ok=True)
    fp = os.path.join(outdir, f"{_safe(nome)}_{_dt.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    wb = Workbook(); ws = wb.active; ws.title = "Dati"
    cols = list(df.columns); ncol = len(cols)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    tc = ws.cell(row=1, column=1, value=nome); tc.font = Font(size=14, bold=True, color="FFFFFF")
    tc.fill = PatternFill("solid", fgColor="1F4E78"); tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    head = 3
    for j, c in enumerate(cols, start=1):
        x = ws.cell(row=head, column=j, value=str(c)); x.font = st["hf"]; x.fill = st["hfill"]; x.alignment = st["ctr"]; x.border = st["bd"]
    vidx = cols.index("VALORE") + 1
    for i, rec in enumerate(df.itertuples(index=False)):
        rr = head + 1 + i
        for j, val in enumerate(rec, start=1):
            x = ws.cell(row=rr, column=j, value=val); x.border = st["bd"]
            if i % 2 == 1: x.fill = st["alt"]
            if j == vidx and isinstance(val, (int, float)):
                x.number_format = "#,##0.###"; x.alignment = st["rgt"]
            else:
                x.alignment = st["lft"]
    for j, c in enumerate(cols, start=1):
        w = max(len(str(c)), int(df[c].astype(str).head(200).map(len).max()) if len(df) else 10)
        ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 10), 45)
    ws.auto_filter.ref = f"A{head}:{get_column_letter(ncol)}{head + len(df)}"
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    _meta(wb, nome, df); wb.save(fp); return fp


def export_matrix(df, nome, outdir, time_label, righe):
    st = _st()
    dims = [c for c in df.columns if c not in (time_label, "VALORE")]
    if righe and righe in df.columns:
        row_col = righe
    else:
        cand = [c for c in dims if c.upper() in ("REF_AREA", "TERRITORIO", "TERRITORY")]
        row_col = cand[0] if cand else max(dims, key=lambda c: df[c].nunique())
    free = [c for c in dims if c != row_col and df[c].nunique() > 1]
    ng = df.groupby(free).ngroups if free else 1
    if ng > 200:
        raise SdmxError(f"Troppe combinazioni per la matrice ({ng} fogli): usa --layout long.")
    years = sorted(df[time_label].dropna().unique(), key=lambda x: str(x))
    wb = Workbook(); wb.remove(wb.active); used = set()
    for gkey, sub in (df.groupby(free, sort=False) if free else [((), df)]):
        gv = (gkey,) if not isinstance(gkey, tuple) else gkey
        gmap = dict(zip(free, gv)) if free else {}
        name = " - ".join(str(gmap[c]) for c in free) if free else "Dati"
        name = re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "Dati"; b, i = name, 1
        while name in used:
            s = f"_{i}"; name = b[:31 - len(s)] + s; i += 1
        used.add(name)
        ws = wb.create_sheet(name); ncol = 1 + len(years)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws.cell(row=1, column=1, value=nome).font = Font(size=13, bold=True, color="1F4E78")
        r = 2
        for c in [c for c in dims if c != row_col]:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
            ws.cell(row=r, column=1, value=f"{c}: {gmap.get(c, sub[c].iloc[0])}").font = Font(size=10, bold=True, color="404040"); r += 1
        r += 1; head = r
        hc = ws.cell(row=head, column=1, value=time_label); hc.font = st["hf"]; hc.fill = st["hfill"]; hc.alignment = st["ctr"]; hc.border = st["bd"]
        for j, y in enumerate(years, start=2):
            yc = ws.cell(row=head, column=j, value=y); yc.font = st["hf"]; yc.fill = st["hfill"]; yc.alignment = st["ctr"]; yc.border = st["bd"]
        piv = sub.pivot_table(index=row_col, columns=time_label, values="VALORE", aggfunc="first")
        ds = head + 1
        for i2, area in enumerate(piv.index):
            rr = ds + i2
            nc = ws.cell(row=rr, column=1, value=str(area)); nc.font = Font(bold=True); nc.border = st["bd"]
            if i2 % 2 == 1: nc.fill = st["alt"]
            for j, y in enumerate(years, start=2):
                v = piv.at[area, y] if y in piv.columns else None
                x = ws.cell(row=rr, column=j, value=None if pd.isna(v) else float(v))
                x.border = st["bd"]; x.alignment = st["rgt"]; x.number_format = "#,##0.###"
                if i2 % 2 == 1: x.fill = st["alt"]
        ws.auto_filter.ref = f"A{head}:{get_column_letter(ncol)}{ds + len(piv.index) - 1}"
        ws.freeze_panes = ws.cell(row=ds, column=2); ws.column_dimensions["A"].width = 34
        for j in range(2, ncol + 1):
            ws.column_dimensions[get_column_letter(j)].width = 14
    _meta(wb, nome, df)
    os.makedirs(outdir, exist_ok=True)
    fp = os.path.join(outdir, f"{_safe(nome)}_{_dt.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    wb.save(fp); return fp



# ===========================================================================
#  API DI ALTO LIVELLO
# ===========================================================================
def export_matrice_unica(df, nome, outdir, time_label):
    import os
    os.makedirs(outdir, exist_ok=True)
    fp = os.path.join(outdir, f"{_safe(nome)}_{_dt.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    
    cand = [c for c in df.columns if "CPA" in c.upper() or "ATECO" in c.upper() or "PARTNER" in c.upper() or "COUNTRY" in c.upper()]
    row_col = cand[0] if cand else [c for c in df.columns if c not in (time_label, "VALORE")][0]
    
    piv = df.pivot_table(index=row_col, columns=time_label, values="VALORE", aggfunc="sum")
    piv.to_excel(fp, sheet_name="Matrice_Unica")
    return fp

def export_matrice_top20(df, nome, outdir, time_label):
    import os
    import pandas as pd
    os.makedirs(outdir, exist_ok=True)
    fp = os.path.join(outdir, f"{_safe(nome)}_{_dt.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    
    cand = [c for c in df.columns if "PARTNER" in c.upper() or "COUNTRY" in c.upper() or "PAESE" in c.upper()]
    row_col = cand[0] if cand else [c for c in df.columns if c not in (time_label, "VALORE")][0]
    
    piv = df.pivot_table(index=row_col, columns=time_label, values="VALORE", aggfunc="sum")
    ultima_colonna = piv.columns[-1]
    piv = piv.sort_values(by=ultima_colonna, ascending=False)
    
    continenti_nomi = ["Africa", "America", "Americhe", "Americas", "Asia", "Europa", "Europe", "Oceania", "AFR", "AME", "ASI", "EUR", "OCE"]
    mask_continenti = piv.index.isin(continenti_nomi)
    df_finale = pd.concat([piv[mask_continenti], piv[~mask_continenti].head(20)])
    
    df_finale.to_excel(fp, sheet_name="Top20_e_Continenti")
    return fp

def export_matrice_solopaesi(df, nome, outdir, time_label):
    import os
    import pandas as pd
    os.makedirs(outdir, exist_ok=True)
    fp = os.path.join(outdir, f"{_safe(nome)}_{_dt.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    
    # Trova la colonna dei paesi
    cand = [c for c in df.columns if "PARTNER" in c.upper() or "COUNTRY" in c.upper() or "PAESE" in c.upper()]
    row_col = cand[0] if cand else [c for c in df.columns if c not in (time_label, "VALORE")][0]
    
    # Crea la matrice base
    piv = df.pivot_table(index=row_col, columns=time_label, values="VALORE", aggfunc="sum")
    
    # Individua dinamicamente l'ultima colonna temporale disponibile
    ultima_colonna = piv.columns[-1]
    
    # Lista dei continenti da escludere tassativamente dal foglio
    continenti_nomi = ["Africa", "America", "Americhe", "Americas", "Asia", "Europa", "Europe", "Oceania", "AFR", "AME", "ASI", "EUR", "OCE"]
    
    # Filtra via i continenti e tiene solo i paesi
    df_paesi = piv[~piv.index.isin(continenti_nomi)]
    
    # Ordina tutti i paesi dal più grande al più piccolo in base all'ultimo trimestre uscito
    df_paesi_ordinati = df_paesi.sort_values(by=ultima_colonna, ascending=False)
    
    # Salva il file risultante
    df_paesi_ordinati.to_excel(fp, sheet_name="Paesi_Ordinati")
    return fp

def esporta(df, time_label, nome, layout, righe, outdir):
    # La nuova parola magica SOLOPAESI
    if layout == "istat" and righe == "SOLOPAESI":
        return export_matrice_solopaesi(df, nome, outdir, time_label)
    # La parola magica TOP20
    if layout == "istat" and righe == "TOP20":
        return export_matrice_top20(df, nome, outdir, time_label)
    # La parola magica UNICA
    if layout == "istat" and righe == "UNICA":
        return export_matrice_unica(df, nome, outdir, time_label)
    if layout == "istat":
        return export_matrix(df, nome, outdir, time_label, righe)
    return export_long(df, nome, outdir, time_label)


def converti_bytes(content, struttura=None, layout="long", righe=None, outdir="output", lang="it", nome="dati"):
    df, time_label = sniff_e_parse(content, struttura, lang)
    return esporta(df, time_label, nome, layout, righe, outdir), len(df)


def converti(dati, struttura=None, layout="long", righe=None, outdir="output", lang="it", nome=None):
    nome = nome or os.path.splitext(os.path.basename(dati))[0]
    if not os.path.exists(dati):
        raise SdmxError(f"File non trovato: {dati}")
    with open(dati, "rb") as fh:
        content = fh.read()
    return converti_bytes(content, struttura, layout, righe, outdir, lang, nome)


def main(argv=None):
    p = argparse.ArgumentParser(description="Converte un file SDMX (XML/JSON/CSV) in Excel.",
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("dati", help="File SDMX (xml/json/csv).")
    p.add_argument("--struttura", help="File SDMX struttura (etichette per XML a codici).")
    p.add_argument("--layout", choices=["long", "istat"], default="long")
    p.add_argument("--righe", help="(layout istat) dimensione nelle righe, es. REF_AREA.")
    p.add_argument("--outdir", default="output")
    p.add_argument("--lang", choices=["it", "en"], default="it")
    a = p.parse_args(argv)
    try:
        path, n = converti(a.dati, a.struttura, a.layout, a.righe, a.outdir, a.lang)
        print(f"OK: {n} osservazioni -> {path}")
    except SdmxError as exc:
        print(f"ERRORE: {exc}", file=sys.stderr); return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
