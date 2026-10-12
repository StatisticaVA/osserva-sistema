#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
report_import_export.py - Report Import/Export (provincia, regione, Italia) da ISTAT Coeweb.
=========================================================================================
Scelti ANNO e TRIMESTRE, scarica i dati ISTAT e crea le tabelle Excel del report:

  Import_export_Sintesi.xlsx            Import_export_Territori.xlsx
  Import_export_Tempo_Trimestrale.xlsx  Import_export_Tempo_Annuale.xlsx
  Import_export_Settori_Export.xlsx     Import_export_Settori_Import.xlsx
  Import_export_Paesi_Export.xlsx       Import_export_Paesi_Import.xlsx
  Export_paesi_xgrafico.xlsx            Import_paesi_xgrafico.xlsx

VALORI CUMULATI: il periodo "anno T2" = T1+T2 dell'anno; lo "stesso periodo
precedente" = T1+T2 dell'anno prima. Con T4 si usa il dato annuale ISTAT
(se pubblicato), altrimenti la somma dei quattro trimestri.
Fanno eccezione, per loro natura:
  - Tempo_Trimestrale: singoli trimestri (ultimi 5 anni + anno scelto);
  - Tempo_Annuale: ultimi 13 anni. Si crea SOLO scegliendo T4 e solo se ISTAT ha
    gia' pubblicato il dato annuale di quell'anno (con T1-T3 non viene creato).

USO:
  python3 report_import_export.py --anno 2026 --trimestre 1 [--outdir CARTELLA]

Richiede sdmx_estrai.py e sdmx_file_to_excel.py nella stessa cartella.
"""
from __future__ import annotations
import argparse, json, os, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import sdmx_estrai as core
import sdmx_file_to_excel as fc

SdmxError = core.SdmxError

# ============================================================================
#  CONFIGURAZIONE
# ============================================================================
_COEWEB = "https://esploradati.istat.it/Coeweb/databrowserhub/api/core/nodes/1/datasets/"
URL_CONG = _COEWEB + "IT1,DF_DCSE_CPA_ATECO2007_COE_A_Q_CONJ,1.0/data"    # annuali+trimestrali dal 2020
URL_STORICO = _COEWEB + "IT1,DF_CPA_AT2007_COE_OPT_A_REG_ALL,1.0/data"  # annuali dal 1991
URL_GEONOM = "https://esploradati.istat.it/SDMXWS/rest/codelist/IT1/CL_COEWEB_GEONOM/1.0"

PROVINCIA, REGIONE, ITALIA = "ITC41", "ITC4", "ITTOT"
NOMI = {PROVINCIA: "Varese", REGIONE: "Lombardia", ITALIA: "Italia"}

# Province della regione, nell'ordine delle tabelle
PROVINCE_REGIONE = [
    ("ITC41", "Varese"), ("ITC42", "Como"), ("ITC44", "Sondrio"), ("ITC45", "Milano"),
    ("ITC46", "Bergamo"), ("ITC47", "Brescia"), ("ITC48", "Pavia"), ("ITC4A", "Cremona"),
    ("ITC4B", "Mantova"), ("ITC43", "Lecco"), ("ITC49", "Lodi"), ("IT108", "Monza e della Brianza"),
]

# Sottosettori manifatturieri (CPA Ateco 2007)
SETTORI = [
    ("CA", "Prodotti alimentari, bevande e tabacco"),
    ("CB", "Prodotti tessili, abbigliamento, pelli accessori"),
    ("CC", "Legno e prodotti in legno; carta e stampa"),
    ("CD", "Coke e prodotti petroliferi raffinati"),
    ("CE", "Sostanze e prodotti chimici"),
    ("CF", "Articoli farmaceutici, chimico-medicinali e botanici"),
    ("CG", "Articoli in gomma e materie plastiche, altri prodotti della lavorazione di minerali non metalliferi"),
    ("CH", "Metalli di base e prodotti in metallo, esclusi macchine e impianti"),
    ("CI", "Computer, apparecchi elettronici e ottici"),
    ("CJ", "Apparecchi elettrici"),
    ("CK", "Macchinari ed apparecchi n.c.a."),
    ("CL", "Mezzi di trasporto"),
    ("CM", "Prodotti altre attività manifatturiere"),
]

CONTINENTI = [("EUR", "Europa"), ("ASI", "Asia"), ("AME", "America"), ("AFR", "Africa"),
              ("OCE_OCE_OTH", "Oceania ed altri territori")]
# Codici "paese" a due lettere che non sono paesi (provviste di bordo, non determinati, ...)
NON_PAESI = {"QP", "QZ", "QW", "QS", "QY", "QV", "QR", "QU", "QX", "QQ", "QO", "XC", "XL"}

ROMANI = {1: "I", 2: "II", 3: "III", 4: "IV"}


# ============================================================================
#  DOWNLOAD
# ============================================================================
def _filtri(freq, aree, tipi, cpa, paesi, y1, y2):
    return [
        {"id": "FREQ", "filterValues": list(freq)},
        {"id": "REF_AREA", "type": "CodeValues", "filterValues": list(aree)},
        {"id": "DATA_TYPE", "type": "CodeValues", "filterValues": list(tipi)},
        {"id": "ADJUSTMENT", "filterValues": ["N"]},
        {"id": "CPA_ATECO2007_COE", "filterValues": list(cpa)},
        {"id": "PARTNER_COUNTRY", "filterValues": list(paesi)},
        {"id": "TIME_PERIOD", "type": "TimeRange", "filterValues": [],
         "from": f"{y1}-01-01T00:00:00", "to": f"{y2}-12-31T00:00:00"},
    ]


def _scarica(url, descr, **kw):
    """POST coeweb -> {(area, tipo, cpa, paese, periodo): valore}."""
    core._notifica(f"scarico {descr}…")
    content = core._post(url, _filtri(**kw), core._ACCEPTS[0], "it")
    try:
        js = json.loads(content.decode("utf-8-sig") if isinstance(content, bytes) else content)
        df, _ = fc.parse_jsonstat(js, "it", codici=True)
    except (ValueError, KeyError) as exc:
        raise SdmxError(f"risposta ISTAT non leggibile ({descr}): {exc}")
    out = {}
    for area, tipo, cpa, paese, per, val in zip(
            df["REF_AREA_CODE"], df["DATA_TYPE_CODE"], df["CPA_ATECO2007_COE_CODE"],
            df["PARTNER_COUNTRY_CODE"], df["TIME_PERIOD_CODE"], df["VALORE"]):
        if isinstance(val, (int, float)) and not pd.isna(val):
            out[(area, tipo, cpa, paese, str(per))] = float(val)
    return out


def _paesi_codelist():
    """Codici e nomi italiani dei paesi partner (codelist GEONOM)."""
    core._notifica("scarico l'elenco dei paesi…")
    r = core._send("GET", URL_GEONOM, headers=core._headers("application/json", "it"))
    js = json.loads(core._check(r))
    nomi = {}
    for c in js["data"]["codelists"][0]["codes"]:
        nomi[c["id"]] = (c.get("names") or {}).get("it") or c.get("name") or c["id"]
    return nomi


# ============================================================================
#  ACCESSO AI DATI (trimestri, anni, cumulati)
# ============================================================================
class Dati:
    def __init__(self):
        self.valori = {}      # dati congiunturali (prevalgono) + storici

    def aggiungi(self, d, prevale=True):
        for k, v in d.items():
            if prevale or k not in self.valori:
                self.valori[k] = v

    def get(self, area, tipo, periodo, cpa="ALL", paese="WORLD"):
        return self.valori.get((area, tipo, cpa, paese, str(periodo)))

    def trimestre(self, area, tipo, anno, t, cpa="ALL", paese="WORLD"):
        return self.get(area, tipo, f"{anno}-Q{t}", cpa, paese)

    def annuale(self, area, tipo, anno, cpa="ALL", paese="WORLD"):
        """Solo il dato annuale pubblicato da ISTAT (nessuna somma dei trimestri)."""
        return self.get(area, tipo, anno, cpa, paese)

    def cumulato(self, area, tipo, anno, trim, cpa="ALL", paese="WORLD"):
        """Somma T1..trim dell'anno (con T4 il dato annuale, se c'e')."""
        if trim == 4:
            v = self.get(area, tipo, anno, cpa, paese)
            if v is not None:
                return v
        qs = [self.trimestre(area, tipo, anno, t, cpa, paese) for t in range(1, trim + 1)]
        if all(x is None for x in qs):
            return None
        return sum(x or 0.0 for x in qs)


def _var(nuovo, vecchio):
    """Variazione % (sul valore assoluto del precedente: vale anche per i saldi)."""
    if nuovo is None or vecchio in (None, 0):
        return None
    return (nuovo - vecchio) / abs(vecchio) * 100


def _quota(parte, totale):
    if parte is None or totale in (None, 0):
        return None
    return parte / totale * 100


def _saldo(e, i):
    return None if e is None or i is None else e - i


# ============================================================================
#  EXCEL
# ============================================================================
def _cella(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, float):
        v = round(v, 2)
        return int(v) if v.is_integer() else v
    return v


def _salva(cartella, nome_file, foglio, intestazioni, righe, col_valori=(), col_pct=()):
    wb = Workbook()
    ws = wb.active
    ws.title = foglio
    ws.append(intestazioni)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E75B6")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in righe:
        ws.append([_cella(v) if isinstance(v, float) else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for j, c in enumerate(row, start=1):
            if not isinstance(c.value, (int, float)):
                continue
            if j in col_valori:
                c.number_format = "#,##0"
            elif j in col_pct:
                c.number_format = "0.00"
    for j, h in enumerate(intestazioni, start=1):
        larg = max([len(str(h))] + [len(str(ws.cell(row=i, column=j).value or ""))
                                    for i in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(larg + 2, 10), 60)
    ws.freeze_panes = "A2"
    path = os.path.join(cartella, nome_file)
    wb.save(path)
    return path


def _p(v):
    """Percentuale arrotondata a 2 decimali (None se non calcolabile)."""
    return None if v is None else round(v, 2)


# ============================================================================
#  TABELLE
# ============================================================================
def tab_sintesi(D, anno, trim, cartella):
    righe, n = [], 0
    for area in (PROVINCIA, REGIONE, ITALIA):
        e1, e0 = D.cumulato(area, "EV", anno, trim), D.cumulato(area, "EV", anno - 1, trim)
        i1, i0 = D.cumulato(area, "IV", anno, trim), D.cumulato(area, "IV", anno - 1, trim)
        for voce, v0, v1, var in (("Export", e0, e1, _p(_var(e1, e0))),
                                  ("Import", i0, i1, _p(_var(i1, i0))),
                                  ("Saldo", _saldo(e0, i0), _saldo(e1, i1), None)):
            n += 1
            righe.append([n, NOMI[area], voce, v0, v1, var])
    return _salva(cartella, "Import_export_Sintesi.xlsx", "Foglio2",
                  ["Ordine", "Territorio", "Bilancia commerciale", "Stesso periodo precedente",
                   "Periodo recente", "Variazione %"], righe, col_valori=(4, 5), col_pct=(6,))


def tab_territori(D, anno, trim, cartella):
    reg = NOMI[REGIONE]
    e_reg, i_reg = D.cumulato(REGIONE, "EV", anno, trim), D.cumulato(REGIONE, "IV", anno, trim)
    righe = []
    for n, (area, nome) in enumerate(PROVINCE_REGIONE + [(REGIONE, reg)], start=1):
        e1, e0 = D.cumulato(area, "EV", anno, trim), D.cumulato(area, "EV", anno - 1, trim)
        i1, i0 = D.cumulato(area, "IV", anno, trim), D.cumulato(area, "IV", anno - 1, trim)
        righe.append([n, nome, e1, _p(_var(e1, e0)), i1, _p(_var(i1, i0)), _saldo(e1, i1),
                      _p(_quota(e1, e_reg)), _p(_quota(i1, i_reg))])
    return _salva(cartella, "Import_export_Territori.xlsx", "Foglio1",
                  ["Ordine", "Territori", "Export", "T.var % tend export", "Import",
                   "T.var % tend import", "Saldo", f"Peso % export su {reg}",
                   f"Peso % import su {reg}"],
                  righe, col_valori=(3, 5, 7), col_pct=(4, 6, 8, 9))


def tab_trimestrale(D, anno, trim, cartella):
    def prec(y, t):
        return (y, t - 1) if t > 1 else (y - 1, 4)

    def serie(tipo, y, t):
        if tipo == "S":
            return _saldo(D.trimestre(PROVINCIA, "EV", y, t), D.trimestre(PROVINCIA, "IV", y, t))
        return D.trimestre(PROVINCIA, tipo, y, t)

    righe = []
    y, t = anno - 5, 1
    while (y, t) <= (anno, trim):
        riga = [y, f"{ROMANI[t]} trim '{y % 100:02d}"]
        for tipo in ("EV", "IV", "S"):
            v = serie(tipo, y, t)
            riga += [v, _p(_var(v, serie(tipo, *prec(y, t)))), _p(_var(v, serie(tipo, y - 1, t)))]
        righe.append(riga)
        y, t = (y, t + 1) if t < 4 else (y + 1, 1)
    return _salva(cartella, "Import_export_Tempo_Trimestrale.xlsx", "Foglio1",
                  ["Anno", "Trimestre", "Export", "T.var% cong. export", "T.var% tend. export",
                   "Import", "T.var% cong. import", "T.var% tend. import", "Saldo",
                   "T.var% cong. saldo", "T.var% tend. saldo"],
                  righe, col_valori=(3, 6, 9), col_pct=(4, 5, 7, 8, 10, 11))


def ultimo_anno_completo(D, anno, trim, area=PROVINCIA):
    y = anno if trim == 4 else anno - 1
    while y > anno - 5 and D.annuale(area, "EV", y) is None:
        y -= 1
    return y


def tab_annuale(D, anno, trim, cartella):
    y_fine = ultimo_anno_completo(D, anno, trim)
    righe = []
    for y in range(y_fine - 12, y_fine + 1):
        e1, e0 = D.annuale(PROVINCIA, "EV", y), D.annuale(PROVINCIA, "EV", y - 1)
        i1, i0 = D.annuale(PROVINCIA, "IV", y), D.annuale(PROVINCIA, "IV", y - 1)
        s1, s0 = _saldo(e1, i1), _saldo(e0, i0)
        righe.append([y, e1, _p(_var(e1, e0)), i1, _p(_var(i1, i0)), s1, _p(_var(s1, s0))])
    return _salva(cartella, "Import_export_Tempo_Annuale.xlsx", "Foglio1",
                  ["Anno", "Export", "T.var % su export anno precedente", "Import",
                   "T.var % su import anno precedente", "Saldo", "T.var% sull'anno precedente"],
                  righe, col_valori=(2, 4, 6), col_pct=(3, 5, 7))


def tab_settori(D, anno, trim, cartella, tipo):
    flusso = "export" if tipo == "EV" else "import"
    cur = lambda c: D.cumulato(PROVINCIA, tipo, anno, trim, cpa=c)
    pre = lambda c: D.cumulato(PROVINCIA, tipo, anno - 1, trim, cpa=c)
    tot_c, tot = cur("C"), cur("ALL")
    righe = []
    for n, (cod, nome) in enumerate(SETTORI, start=1):
        v = cur(cod)
        righe.append([n, nome, v, _p(_quota(v, tot_c)), _p(_var(v, pre(cod)))])
    n = len(SETTORI)
    righe.append([n + 1, "Totale manifatturiero", tot_c, 100.0 if tot_c else None,
                  _p(_var(tot_c, pre("C")))])
    righe.append([n + 2, f"% manifatturiero sul totale dell'{flusso}", tot,
                  _p(_quota(tot_c, tot)), _p(_var(tot, pre("ALL")))])
    nome_file = f"Import_export_Settori_{flusso.capitalize()}.xlsx"
    return _salva(cartella, nome_file, "Foglio1",
                  ["Ordine", "Dettaglio manifatturiero", "Valori in euro",
                   f"% {flusso} sul totale manifatturiero", "T.var% tend."],
                  righe, col_valori=(3,), col_pct=(4, 5))


def tab_paesi(D, anno, trim, cartella, tipo, nomi_paesi):
    flusso = "Export" if tipo == "EV" else "Import"
    cur = lambda p: D.cumulato(PROVINCIA, tipo, anno, trim, paese=p)
    pre = lambda p: D.cumulato(PROVINCIA, tipo, anno - 1, trim, paese=p)
    mondo = cur("WORLD")
    riga = lambda n, nome, p: [n, nome, cur(p), _p(_quota(cur(p), mondo)), _p(_var(cur(p), pre(p)))]

    continenti = sorted(CONTINENTI, key=lambda c: cur(c[0]) or 0, reverse=True)
    righe = [riga(n, nome, cod) for n, (cod, nome) in enumerate(continenti, start=1)]
    grafico = [list(r) for r in righe]
    n = len(righe)
    righe.append([n + 1, "MONDO", mondo, 100.0 if mondo else None, _p(_var(mondo, pre("WORLD")))])
    righe.append([n + 2, "Primi 20 paesi partner", None, None, None])
    paesi = [p for p in nomi_paesi if len(p) == 2 and p not in NON_PAESI and cur(p)]
    paesi.sort(key=lambda p: cur(p), reverse=True)
    for k, p in enumerate(paesi[:20], start=n + 3):
        righe.append(riga(k, nomi_paesi[p], p))

    intest = ["Ordine", "Aree geografiche", flusso, "Composizione %", "T.var% tend."]
    p1 = _salva(cartella, f"Import_export_Paesi_{flusso}.xlsx", "Foglio1", intest, righe,
                col_valori=(3,), col_pct=(4, 5))
    p2 = _salva(cartella, f"{flusso}_paesi_xgrafico.xlsx", "Foglio1", intest, grafico,
                col_valori=(3,), col_pct=(4, 5))
    return [p1, p2]


# ============================================================================
#  REPORT COMPLETO
# ============================================================================
def genera_report(anno, trim, outdir, log=None):
    """Scarica i dati e crea tutte le tabelle. Ritorna (cartella, [file creati]).
    log: callable(testo, tag) opzionale per i messaggi (tag: 'ok' / 'err')."""
    anno, trim = int(anno), int(trim)
    if trim not in (1, 2, 3, 4):
        raise SdmxError("Il trimestre deve essere 1, 2, 3 o 4.")
    scrivi = log or (lambda t, tag=None: print(t))
    settori = [c for c, _ in SETTORI]
    lombardia = [c for c, _ in PROVINCE_REGIONE]
    D = Dati()

    # 1) totali trimestrali/annuali di Italia, regione e province (anno-6 .. anno)
    D.aggiungi(_scarica(URL_CONG, "totali di Italia, regione e province",
                        freq=["A", "Q"], aree=[ITALIA, REGIONE] + lombardia, tipi=["EV", "IV"],
                        cpa=["ALL"], paesi=["WORLD"], y1=anno - 6, y2=anno))
    if D.trimestre(PROVINCIA, "EV", anno, trim) is None:
        ultimi = sorted({k[4] for k in D.valori if "-Q" in k[4] and k[0] == PROVINCIA})
        raise SdmxError(f"ISTAT non ha ancora pubblicato i dati del {trim}° trimestre {anno}"
                        + (f" (ultimo disponibile: {ultimi[-1]})." if ultimi else "."))

    cartella = os.path.join(outdir, f"Report_import_export_{anno}_T{trim}")
    os.makedirs(cartella, exist_ok=True)
    creati = []

    def fai(nome, fn, *args):
        try:
            res = fn(*args)
            for p in (res if isinstance(res, list) else [res]):
                creati.append(p)
                scrivi(os.path.basename(p), "ok")
        except SdmxError as exc:
            if "annullata" in str(exc):
                raise
            scrivi(f"{nome}: {exc}", "err")
        except Exception as exc:
            scrivi(f"{nome}: errore {exc}", "err")

    fai("Sintesi", tab_sintesi, D, anno, trim, cartella)
    fai("Territori", tab_territori, D, anno, trim, cartella)
    fai("Tempo trimestrale", tab_trimestrale, D, anno, trim, cartella)

    # 2) Tempo_Annuale: solo con T4 e solo se ISTAT ha pubblicato il dato annuale
    if trim != 4:
        scrivi("Tempo annuale: si genera solo scegliendo il T4", "info")
    elif D.annuale(PROVINCIA, "EV", anno) is None or D.annuale(PROVINCIA, "IV", anno) is None:
        scrivi(f"Tempo annuale: dato annuale {anno} non ancora pubblicato da ISTAT", "info")
    else:
        try:
            D.aggiungi(_scarica(URL_STORICO, "serie storica annuale",
                                freq=["A"], aree=[PROVINCIA], tipi=["EV", "IV"],
                                cpa=["ALL"], paesi=["WORLD"], y1=anno - 14, y2=anno),
                       prevale=False)
            fai("Tempo annuale", tab_annuale, D, anno, trim, cartella)
        except SdmxError as exc:
            if "annullata" in str(exc):
                raise
            scrivi(f"Tempo annuale: serie storica non disponibile: {exc}", "err")

    # 3) settori manifatturieri della provincia
    try:
        D.aggiungi(_scarica(URL_CONG, "settori manifatturieri",
                            freq=["A", "Q"], aree=[PROVINCIA], tipi=["EV", "IV"],
                            cpa=["ALL", "C"] + settori, paesi=["WORLD"], y1=anno - 1, y2=anno))
        fai("Settori export", tab_settori, D, anno, trim, cartella, "EV")
        fai("Settori import", tab_settori, D, anno, trim, cartella, "IV")
    except SdmxError as exc:
        if "annullata" in str(exc):
            raise
        scrivi(f"Settori: {exc}", "err")

    # 4) paesi e continenti della provincia
    try:
        nomi_paesi = _paesi_codelist()
        nomi_paesi.update(dict(CONTINENTI))
        codici = [p for p in nomi_paesi if len(p) == 2 and p not in NON_PAESI]
        D.aggiungi(_scarica(URL_CONG, "paesi e continenti",
                            freq=["A", "Q"], aree=[PROVINCIA], tipi=["EV", "IV"], cpa=["ALL"],
                            paesi=codici + [c for c, _ in CONTINENTI] + ["WORLD"],
                            y1=anno - 1, y2=anno))
        fai("Paesi export", tab_paesi, D, anno, trim, cartella, "EV", nomi_paesi)
        fai("Paesi import", tab_paesi, D, anno, trim, cartella, "IV", nomi_paesi)
    except SdmxError as exc:
        if "annullata" in str(exc):
            raise
        scrivi(f"Paesi: {exc}", "err")

    return cartella, creati


def main(argv=None):
    p = argparse.ArgumentParser(description="Report Import/Export ISTAT (valori cumulati).")
    p.add_argument("--anno", type=int, required=True, help="es. 2026")
    p.add_argument("--trimestre", type=int, required=True, choices=[1, 2, 3, 4])
    p.add_argument("--outdir", default="output")
    a = p.parse_args(argv)
    core.PROGRESSO = lambda t: print(f"  … {t}", flush=True)
    try:
        cartella, creati = genera_report(a.anno, a.trimestre, a.outdir,
                                         log=lambda t, tag=None: print({"ok": "  OK ", "info": "  INFO "}.get(tag, "  ERRORE ") + t, flush=True))
    except SdmxError as exc:
        print(f"ERRORE: {exc}", file=sys.stderr)
        return 1
    print(f"\nCompletato: {len(creati)} file in {cartella}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
